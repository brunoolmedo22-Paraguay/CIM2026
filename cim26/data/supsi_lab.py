"""Importador reprodutível da caracterização de laboratório SUPSI (IEC 61853-1/-2).

Lê ``IEC61853_1-2.xlsx`` (folha ``61853-1 & 2``) SEM escrever sobre o original.
Matrizes: Pmp linhas 9–15, Isc 19–25, Voc 29–35; colunas B..E = 15, 25, 50, 75 °C;
coluna A = irradiância 1100→100 W/m²; coluna F = incerteza relativa u (k=2).

Auditoria do «32»: no XML do arquivo, as células sem medição são do tipo
string compartilhada (``t="s"``) com ``<v>32</v>``, isto é, um ÍNDICE para a
tabela ``sharedStrings`` cujo item 32 é a cadeia vazia. Um leitor XML ingênuo
veria «32»; o openpyxl devolve ''. Em ambos os casos a célula é ausência de
medição. Este módulo trata como ausente toda célula não numérica e registra a
auditoria XML em ``audit_missing_cells``.
"""

from __future__ import annotations

import hashlib
import json

import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

import pandas as pd

from ..paths import SUPSI_LAB_LONG_CSV, SUPSI_MODULE_JSON, SUPSI_RAW_LAB_XLSX
from ..pvcore.module import KIND_EXPERIMENTAL, ModuleSpec

SHEET = "61853-1 & 2"
BLOCKS = {"Pmp": (9, 15, "W"), "Isc": (19, 25, "A"), "Voc": (29, 35, "V")}
TEMP_COLS = {"B": 15.0, "C": 25.0, "D": 50.0, "E": 75.0}
# Condições exigidas pela IEC 61853-1 (22 pares G–T) — usadas como verificação cruzada
IEC61853_1_CONDITIONS = {
    1100: (25, 50, 75), 1000: (15, 25, 50, 75), 800: (15, 25, 50, 75),
    600: (15, 25, 50, 75), 400: (15, 25, 50), 200: (15, 25), 100: (15, 25),
}
_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


@dataclass
class LabCharacterization:
    long: pd.DataFrame                       # variable, irradiance_poa_Wm2, temperature_module_C, measured_value, unit, xlsx_row, xlsx_col, u_rel_k2
    coefficients: dict                       # α, β, δ (%/°C) e incertezas relativas
    thermal: dict                            # u0, u1, NMOT
    header: dict                             # tecnologia, laboratório, contato
    audit_missing_cells: list = field(default_factory=list)
    source_sha256: str = ""
    source_path: str = ""

    def matrix(self, variable: str) -> pd.DataFrame:
        d = self.long[self.long["variable"] == variable]
        return d.pivot(index="irradiance_poa_Wm2", columns="temperature_module_C",
                       values="measured_value").sort_index(ascending=False)

    @property
    def stc(self) -> dict:
        d = self.long[(self.long.irradiance_poa_Wm2 == 1000) & (self.long.temperature_module_C == 25)]
        return {r.variable: float(r.measured_value) for r in d.itertuples()}


def _audit_xml(path: Path, cells: list[str]) -> list[dict]:
    """Inspeciona o XML bruto das células indicadas (tipo, valor bruto, string resolvida)."""
    out = []
    with zipfile.ZipFile(path) as z:
        shared = []
        if "xl/sharedStrings.xml" in z.namelist():
            root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for si in root.findall("m:si", _NS):
                shared.append("".join(t.text or "" for t in si.iter(f"{{{_NS['m']}}}t")))
        sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
        index = {c.get("r"): c for c in sheet.iter(f"{{{_NS['m']}}}c")}
    for ref in cells:
        c = index.get(ref)
        if c is None:
            out.append({"cell": ref, "xml_type": None, "xml_raw_value": None, "resolved": None})
            continue
        v = c.find("m:v", _NS)
        raw = None if v is None else v.text
        typ = c.get("t")
        resolved = shared[int(raw)] if (typ == "s" and raw is not None and raw.isdigit()
                                        and int(raw) < len(shared)) else raw
        out.append({"cell": ref, "xml_type": typ, "xml_raw_value": raw, "resolved": resolved,
                     "interpretation": "ausência de medição (índice de string vazia)"
                     if (typ == "s" and resolved == "") else "verificar"})
    return out


def read_iec61853_xlsx(path: Path | str = SUPSI_RAW_LAB_XLSX) -> LabCharacterization:
    from openpyxl import load_workbook  # dependência só da importação, não da simulação

    path = Path(path)
    wb = load_workbook(path, data_only=True, read_only=False)
    ws = wb[SHEET]
    rows, missing = [], []
    for var, (r0, r1, unit) in BLOCKS.items():
        header_row = r0 - 1
        for col, t_nom in TEMP_COLS.items():
            hdr = ws[f"{col}{header_row}"].value
            if float(hdr) != t_nom:
                raise ValueError(f"Cabeçalho inesperado {col}{header_row}={hdr!r} (esperado {t_nom}).")
        for r in range(r0, r1 + 1):
            g = ws[f"A{r}"].value
            u_rel = ws[f"F{r}"].value
            for col, t_nom in TEMP_COLS.items():
                v = ws[f"{col}{r}"].value
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    rows.append(dict(variable=var, irradiance_poa_Wm2=float(g), temperature_module_C=t_nom,
                                     measured_value=float(v), unit=unit, xlsx_row=r, xlsx_col=col,
                                     u_rel_k2=float(u_rel) if isinstance(u_rel, (int, float)) else None))
                else:
                    missing.append(f"{col}{r}")
    long = pd.DataFrame(rows)
    # Coeficientes (linha 9) e incertezas relativas (linha 10)
    coeffs = {
        "alpha_isc_pct_per_C": float(ws["I9"].value), "beta_voc_pct_per_C": float(ws["J9"].value),
        "gamma_pmax_pct_per_C": float(ws["K9"].value),
        "u_rel_k2_alpha": float(ws["I10"].value), "u_rel_k2_beta": float(ws["J10"].value),
        "u_rel_k2_gamma": float(ws["K10"].value),
        "standard": str(ws["I6"].value),
    }
    thermal = {"u0_W_m2K": float(ws["U9"].value), "u1_W_s_m3K": float(ws["V9"].value),
               "nmot_C": float(ws["W9"].value), "standard": str(ws["U6"].value)}
    header = {"technology": ws["B1"].value, "lab": ws["A2"].value, "contact": ws["B3"].value,
              "power_rating_standard": ws["B6"].value,
              "label_quirks": {"B27": ws["B27"].value}}
    wb.close()
    audit = _audit_xml(path, missing)
    return LabCharacterization(long=long, coefficients=coeffs, thermal=thermal, header=header,
                               audit_missing_cells=audit, source_sha256=sha256_file(path),
                               source_path=str(path))


def check_lab(lab: LabCharacterization, module_json: dict | None = None, tol_rel: float = 1e-9) -> list[dict]:
    """Verificações de aceitação: contagens 22×3, padrão IEC, STC × ficha experimental."""
    checks = []
    counts = lab.long.groupby("variable").size().to_dict()
    checks.append({"check": "22 observações por grandeza", "ok": counts == {"Isc": 22, "Pmp": 22, "Voc": 22},
                   "detail": counts})
    pattern_ok = True
    for var in BLOCKS:
        d = lab.long[lab.long.variable == var]
        got = {(int(r.irradiance_poa_Wm2), int(r.temperature_module_C)) for r in d.itertuples()}
        exp = {(g, t) for g, ts in IEC61853_1_CONDITIONS.items() for t in ts}
        pattern_ok &= got == exp
    checks.append({"check": "padrão de posições = matriz IEC 61853-1", "ok": pattern_ok, "detail": ""})
    checks.append({"check": "nenhum valor 32 tratado como medição",
                   "ok": not (lab.long.measured_value == 32).any(), "detail": f"{len(lab.audit_missing_cells)} células ausentes auditadas"})
    sentinel_ok = all(a["xml_type"] == "s" and a["xml_raw_value"] == "32" and a["resolved"] == ""
                      for a in lab.audit_missing_cells)
    checks.append({"check": "células ausentes = string compartilhada nº 32 (vazia)", "ok": sentinel_ok,
                   "detail": ", ".join(a["cell"] for a in lab.audit_missing_cells)})
    if module_json:
        stc = lab.stc
        for k_lab, k_js in (("Pmp", "power_stc_measured_W"), ("Isc", "isc_stc_A"), ("Voc", "voc_stc_V")):
            ok = abs(stc[k_lab] - module_json[k_js]) <= tol_rel * abs(module_json[k_js])
            checks.append({"check": f"STC {k_lab} = ficha experimental", "ok": ok,
                           "detail": f"{stc[k_lab]!r} vs {module_json[k_js]!r}"})
        for k_lab, k_js in (("alpha_isc_pct_per_C", "alpha_isc_pct_per_C"), ("beta_voc_pct_per_C", "beta_voc_pct_per_C"),
                            ("gamma_pmax_pct_per_C", "gamma_pmax_pct_per_C")):
            checks.append({"check": f"{k_js} = ficha", "ok": lab.coefficients[k_lab] == module_json[k_js],
                           "detail": f"{lab.coefficients[k_lab]} vs {module_json[k_js]}"})
        for k_lab, k_js in (("u0_W_m2K", "u0_W_m2K"), ("u1_W_s_m3K", "u1_W_s_m3K"), ("nmot_C", "nmot_C")):
            checks.append({"check": f"{k_js} = ficha", "ok": lab.thermal[k_lab] == module_json[k_js],
                           "detail": f"{lab.thermal[k_lab]} vs {module_json[k_js]}"})
    return checks


def write_long_csv(lab: LabCharacterization, out_path: Path) -> Path:
    """Regenera o derivado ``lab_IEC61853_matrix_long.csv`` (nunca sobre o original)."""
    out_path = Path(out_path)
    if out_path.resolve() == SUPSI_LAB_LONG_CSV.resolve():
        raise PermissionError("Não sobrescrever o derivado declarado do pacote; escolha outro destino.")
    cols = ["variable", "irradiance_poa_Wm2", "temperature_module_C", "measured_value", "unit", "xlsx_row", "xlsx_col"]
    lab.long[cols].to_csv(out_path, index=False, lineterminator="\r\n")
    return out_path


def read_long_csv(path: Path | str = SUPSI_LAB_LONG_CSV) -> pd.DataFrame:
    return pd.read_csv(path)


def load_supsi_module_json(path: Path | str = SUPSI_MODULE_JSON) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def supsi_module_spec(path: Path | str = SUPSI_MODULE_JSON) -> ModuleSpec:
    """Ficha experimental SUPSI → ModuleSpec. Vmp, Imp, Ns, fabricante e modelo permanecem None."""
    js = load_supsi_module_json(path)
    return ModuleSpec(
        id=js["id"], kind=KIND_EXPERIMENTAL, technology=js["technology"],
        p_stc_W=float(js["power_stc_measured_W"]), v_oc_V=float(js["voc_stc_V"]), i_sc_A=float(js["isc_stc_A"]),
        v_mp_V=js["vmp_stc_V"], i_mp_A=js["imp_stc_A"], n_cells_series=js["series_cell_count"],
        alpha_isc_pct_C=float(js["alpha_isc_pct_per_C"]), beta_voc_pct_C=float(js["beta_voc_pct_per_C"]),
        gamma_pmax_pct_C=float(js["gamma_pmax_pct_per_C"]), noct_C=None,
        u0_W_m2K=float(js["u0_W_m2K"]), u1_W_s_m3K=float(js["u1_W_s_m3K"]),
        manufacturer=js["manufacturer"], model=js["commercial_model"],
        source=f"{js['source_original']} | {Path(path).name}",
        p_stc_meaning="Pmp medida em STC (laboratório SUPSI, IEC 61853-1)",
        notes=js.get("uncertainty", ""), extra={"nmot_C": js.get("nmot_C")},
    )


